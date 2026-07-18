#!/usr/bin/env python3
"""Generate PPTX presentation: Вещевая служба ВС РФ."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import nsmap
from pptx.oxml import parse_xml
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUT = ROOT / "Вещевая_служба_ВС_РФ.pptx"

# Widescreen 16:9
W, H = Inches(13.333), Inches(7.5)

BG = RGBColor(0x12, 0x16, 0x0F)
PANEL = RGBColor(0x1B, 0x22, 0x16)
PANEL2 = RGBColor(0x24, 0x30, 0x1F)
OLIVE = RGBColor(0x5F, 0x73, 0x48)
OLIVE_B = RGBColor(0x8F, 0xA8, 0x6A)
GOLD = RGBColor(0xC9, 0xA8, 0x5C)
SAND = RGBColor(0xE6, 0xDC, 0xC4)
MUTED = RGBColor(0xB7, 0xB0, 0x9A)
WHITE = RGBColor(0xF3, 0xEF, 0xE4)
LINE = RGBColor(0x3A, 0x45, 0x30)


def set_run(run, size=14, bold=False, color=WHITE, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, color=BG):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    # send to back
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)
    return shape


def add_rect(slide, l, t, w, h, fill=PANEL, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.color.rgb = LINE
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def add_text(slide, l, t, w, h, text, size=14, bold=False, color=WHITE, font="Calibri", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def add_bullets(slide, l, t, w, h, items, size=13, color=MUTED, bold_first=False):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        p.level = 0
        run = p.add_run()
        run.text = "•  " + item
        set_run(run, size=size, bold=False, color=color)
    return box


def eyebrow(slide, text, l=Inches(0.7), t=Inches(0.35)):
    add_text(slide, l, t, Inches(10), Inches(0.35), text.upper(), size=11, bold=True, color=GOLD, font="Calibri")


def title(slide, text, l=Inches(0.7), t=Inches(0.65), w=Inches(12), size=32):
    add_text(slide, l, t, w, Inches(0.9), text, size=size, bold=True, color=SAND, font="Georgia")


def subtitle(slide, text, l=Inches(0.7), t=Inches(1.4), w=Inches(11.5), size=14):
    add_text(slide, l, t, w, Inches(0.55), text, size=size, color=MUTED)


def footer(slide, n, total):
    add_text(slide, Inches(0.7), Inches(7.1), Inches(8), Inches(0.3),
             "Вещевая служба ВС РФ  ·  справочный обзор", size=10, color=MUTED)
    add_text(slide, Inches(11.2), Inches(7.1), Inches(1.5), Inches(0.3),
             f"{n} / {total}", size=10, color=GOLD, align=PP_ALIGN.RIGHT)


def new_slide(prs):
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    return slide


def panel_card(slide, l, t, w, h, heading, body_lines, heading_size=16):
    add_rect(slide, l, t, w, h, fill=PANEL)
    add_text(slide, l + Inches(0.25), t + Inches(0.18), w - Inches(0.4), Inches(0.4),
             heading, size=heading_size, bold=True, color=GOLD, font="Georgia")
    add_bullets(slide, l + Inches(0.2), t + Inches(0.55), w - Inches(0.35), h - Inches(0.7),
                body_lines, size=12, color=MUTED)


def add_table(slide, l, t, w, h, rows, col_widths=None, header=True):
    n_rows = len(rows)
    n_cols = len(rows[0])
    table_shape = slide.shapes.add_table(n_rows, n_cols, l, t, w, h)
    table = table_shape.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            table.columns[i].width = cw
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT
                for run in p.runs:
                    is_header = header and r == 0
                    set_run(run, size=10 if not is_header else 10, bold=is_header,
                            color=GOLD if is_header else SAND)
            # cell fill
            fill = cell.fill
            fill.solid()
            if header and r == 0:
                fill.fore_color.rgb = PANEL2
            elif r % 2 == 0:
                fill.fore_color.rgb = PANEL
            else:
                fill.fore_color.rgb = RGBColor(0x16, 0x1C, 0x12)
    return table_shape


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    total = 28

    # 1 Title
    s = new_slide(prs)
    if (ASSETS / "hero-bg.png").exists():
        s.shapes.add_picture(str(ASSETS / "hero-bg.png"), 0, 0, W, H)
        # dark overlay
        ov = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(3.2), W, Inches(4.3))
        ov.fill.solid()
        ov.fill.fore_color.rgb = RGBColor(0x12, 0x16, 0x0F)
        ov.line.fill.background()
        # soften by covering lower area only - already dark bg underneath for top? add top fade
        top = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(3.5))
        top.fill.solid()
        top.fill.fore_color.rgb = RGBColor(0x12, 0x16, 0x0F)
        top.fill.fore_color.brightness = -0.2
        # python-pptx brightness may not work well; use solid with transparency via alpha if possible
        top.line.fill.background()
        # Actually solid overlay hides image - better approach: only bottom gradient bar
        # Remove top overlay by deleting - simpler: don't add top, just bottom panel
        spTree = s.shapes._spTree
        spTree.remove(top._element)
    add_text(s, Inches(0.8), Inches(3.6), Inches(11), Inches(0.35),
             "ВООРУЖЁННЫЕ СИЛЫ РОССИЙСКОЙ ФЕДЕРАЦИИ", size=12, bold=True, color=GOLD)
    add_text(s, Inches(0.8), Inches(4.0), Inches(11), Inches(1.2),
             "Вещевая служба", size=48, bold=True, color=SAND, font="Georgia")
    add_text(s, Inches(0.8), Inches(5.3), Inches(10.5), Inches(1.0),
             "Подробный обзор системы вещевого обеспечения: законодательство, приказы, "
             "нормы снабжения, ВКПО, размеры и ростовка, сроки носки, порядок выдачи и компенсации.",
             size=15, color=MUTED)
    footer(s, 1, total)

    # 2 Agenda
    s = new_slide(prs)
    eyebrow(s, "Навигация")
    title(s, "Содержание презентации")
    subtitle(s, "Двадцать восемь слайдов — от правовой базы до практических таблиц размеров и сроков носки.")
    items = [
        ("01", "Миссия и место в системе", "Задачи вещевой службы в тыловом обеспечении"),
        ("02", "Законодательство и приказы", "76-ФЗ, ПП № 390, приказ МО № 500"),
        ("03", "Виды имущества", "Личное, инвентарное, специальное"),
        ("04", "ВКПО", "Состав комплекта и порядок выдачи"),
        ("05", "Сроки носки", "Типовые сроки и климатические поправки"),
        ("06", "Размеры и ростовка", "ГОСТ, замеры, таблицы"),
        ("07", "Специальные нормы", "Холод, ВМФ, ССО, медицина, спорт"),
        ("08", "Практика", "Выдача, учёт, компенсации, изменения 2024–2025"),
    ]
    for i, (num, h, d) in enumerate(items):
        col = i % 2
        row = i // 2
        l = Inches(0.7 + col * 6.2)
        t = Inches(2.1 + row * 1.1)
        add_rect(s, l, t, Inches(5.9), Inches(0.95), fill=PANEL)
        add_text(s, l + Inches(0.2), t + Inches(0.15), Inches(0.7), Inches(0.6), num, size=20, bold=True, color=GOLD, font="Georgia")
        add_text(s, l + Inches(1.0), t + Inches(0.15), Inches(4.6), Inches(0.35), h, size=14, bold=True, color=SAND)
        add_text(s, l + Inches(1.0), t + Inches(0.5), Inches(4.6), Inches(0.35), d, size=11, color=MUTED)
    footer(s, 2, total)

    # 3 Mission
    s = new_slide(prs)
    eyebrow(s, "Раздел 01")
    title(s, "Миссия вещевой службы")
    subtitle(s, "Вещевое обеспечение — часть материального обеспечения войск: форма, обувь, снаряжение, БПО и учёт.")
    cards = [
        ("Главная цель", [
            "Своевременно обеспечивать военнослужащих вещевым имуществом по нормам",
            "Учёт климата, рода войск и характера задач",
            "Планирование, снабжение, учёт, ремонт",
        ]),
        ("Ключевые функции", [
            "Расчёт потребности и заявок по ростовке",
            "Хранение и отпуск со складов",
            "Подгонка, клеймение, замена по срокам",
            "Банно-прачечное обслуживание",
        ]),
        ("Для кого", [
            "Военнослужащие по контракту и по призыву",
            "Курсанты военных вузов",
            "Граждане на военных сборах",
            "Отдельные категории ведомственной охраны МО",
        ]),
    ]
    for i, (h, lines) in enumerate(cards):
        panel_card(s, Inches(0.7 + i * 4.15), Inches(2.15), Inches(3.95), Inches(4.4), h, lines)
    footer(s, 3, total)

    # 4 System flow
    s = new_slide(prs)
    eyebrow(s, "Раздел 01")
    title(s, "Место в системе обеспечения")
    subtitle(s, "От центральных органов до вещевого склада воинской части.")
    steps = [
        ("01", "Центральный уровень", "Нормы, закупки, централизованное снабжение"),
        ("02", "Округ / флот", "Наряды и разнарядки по ростовке"),
        ("03", "ЦМТО / склады", "Хранение запасов, отпуск частям"),
        ("04", "Воинская часть", "Склад, карточки, выдача, подгонка"),
        ("05", "Подразделение", "Закрепление инвентаря, контроль носки"),
    ]
    for i, (n, h, d) in enumerate(steps):
        l = Inches(0.45 + i * 2.55)
        add_rect(s, l, Inches(2.4), Inches(2.4), Inches(3.6), fill=PANEL)
        add_text(s, l + Inches(0.15), Inches(2.6), Inches(2.1), Inches(0.5), n, size=22, bold=True, color=OLIVE_B, font="Georgia")
        add_text(s, l + Inches(0.15), Inches(3.2), Inches(2.1), Inches(0.7), h, size=14, bold=True, color=SAND)
        add_text(s, l + Inches(0.15), Inches(4.0), Inches(2.1), Inches(1.6), d, size=12, color=MUTED)
        if i < 4:
            add_text(s, l + Inches(2.25), Inches(3.9), Inches(0.3), Inches(0.4), "→", size=18, color=GOLD)
    footer(s, 4, total)

    # 5 Legal
    s = new_slide(prs)
    eyebrow(s, "Раздел 02")
    title(s, "Правовая пирамида")
    subtitle(s, "Нормы закреплены на уровне закона, постановлений Правительства и ведомственных приказов.")
    laws = [
        ("76-ФЗ, ст. 14", "О статусе военнослужащих — гарантия вещевого обеспечения и право ряда категорий на компенсацию."),
        ("ПП № 390\n22.06.2006", "О вещевом обеспечении в мирное время. Нормы, правила владения, БПО, компенсации. Изм. ПП № 1116/2025."),
        ("Приказ МО\n№ 500", "14.08.2017 (ред. 27.06.2024) — главный «вещевой» документ ВС РФ: Порядок, нормы, специмущество."),
        ("Расп. № 1014-р\n25.05.2016", "Стоимость предметов для выдачи взамен неполученного и расчётов компенсаций."),
    ]
    for i, (code, desc) in enumerate(laws):
        t = Inches(2.1 + i * 1.1)
        add_rect(s, Inches(0.7), t, Inches(12), Inches(1.0), fill=PANEL)
        add_text(s, Inches(0.95), t + Inches(0.15), Inches(2.4), Inches(0.7), code, size=13, bold=True, color=GOLD, font="Georgia")
        add_text(s, Inches(3.5), t + Inches(0.2), Inches(8.9), Inches(0.65), desc, size=13, color=MUTED)
    footer(s, 5, total)

    # 6 Order 500 structure
    s = new_slide(prs)
    eyebrow(s, "Раздел 02")
    title(s, "Приказ МО РФ № 500 — структура")
    subtitle(s, "«О вещевом обеспечении в Вооружённых Силах Российской Федерации на мирное время».")
    apps = [
        ("Прил. № 1", "Порядок обеспечения: выдача, замена, перевод, увольнение, ВКПО."),
        ("Прил. № 2", "Нормы переходящих и страховых запасов, моющие материалы."),
        ("Прил. № 3", "Нормы содержания имущества на одного военнослужащего / койку."),
        ("Прил. № 4+", "Специальное, спортивное, полевое имущество, палатки, брезенты."),
    ]
    for i, (h, d) in enumerate(apps):
        panel_card(s, Inches(0.7 + i * 3.15), Inches(2.15), Inches(3.0), Inches(2.8), h, [d], heading_size=15)
    stats = [("45+", "спецнорм"), ("3", "категории качества"), ("10%", "запас на призыве"), ("ВКПО", "базовый комплект")]
    for i, (a, b) in enumerate(stats):
        l = Inches(0.7 + i * 3.15)
        add_rect(s, l, Inches(5.2), Inches(3.0), Inches(1.3), fill=PANEL2)
        add_text(s, l + Inches(0.2), Inches(5.35), Inches(2.6), Inches(0.5), a, size=22, bold=True, color=OLIVE_B, font="Georgia")
        add_text(s, l + Inches(0.2), Inches(5.9), Inches(2.6), Inches(0.4), b, size=12, color=MUTED)
    footer(s, 6, total)

    # 7 Property types
    s = new_slide(prs)
    eyebrow(s, "Раздел 03")
    title(s, "Виды вещевого имущества")
    subtitle(s, "От вида зависят право собственности, обязанность сдачи и порядок замены.")
    cards = [
        ("Личное пользование", [
            "Форма, бельё, обувь повседневного назначения",
            "Учёт в карточке вещевого довольствия",
            "Замена по срокам носки",
            "Для отдельных категорий — компенсация деньгами",
        ]),
        ("Инвентарное", [
            "ВКПО (кроме исключений), палатки, постовая одежда",
            "Не переходит в личную собственность",
            "Сдаётся при переводе / увольнении",
            "Замена — после сдачи изношенного",
        ]),
        ("Специальное", [
            "Лётное, техническое, климатическое, ВМФ, ССО",
            "Нормы № 1–45 и др.",
            "Часто инвентарное",
            "Выдача по назначению / должности",
        ]),
    ]
    for i, (h, lines) in enumerate(cards):
        panel_card(s, Inches(0.7 + i * 4.15), Inches(2.15), Inches(3.95), Inches(4.4), h, lines)
    footer(s, 7, total)

    # 8 Categories
    s = new_slide(prs)
    eyebrow(s, "Раздел 03")
    title(s, "Категории качества имущества")
    cards = [
        ("I категория", ["Новое имущество, не бывшее в употреблении.", "Офицеры и прапорщики — ВКПО I категории."]),
        ("II категория", ["Бывшее в употреблении, срок не истёк;", "либо выслужившее срок, но годное к использованию.", "Типично для призывников и периода испытания."]),
        ("III категория", ["Непригодное к использованию, срок истёк — списание.", "Расходники (носки, мыло) на категории не делятся."]),
    ]
    for i, (h, lines) in enumerate(cards):
        panel_card(s, Inches(0.7 + i * 4.15), Inches(2.15), Inches(3.95), Inches(4.4), h, lines)
    footer(s, 8, total)

    # 9 VKPO
    s = new_slide(prs)
    eyebrow(s, "Раздел 04")
    title(s, "ВКПО — всесезонный комплект")
    subtitle(s, "Базовый полевой комплект. Предметы ВКПО — преимущественно инвентарные.")
    if (ASSETS / "vkpo-kit.png").exists():
        s.shapes.add_picture(str(ASSETS / "vkpo-kit.png"), Inches(0.7), Inches(2.15), Inches(6.2), Inches(4.4))
    add_rect(s, Inches(7.2), Inches(2.15), Inches(5.4), Inches(4.4), fill=PANEL)
    add_text(s, Inches(7.45), Inches(2.35), Inches(5.0), Inches(0.4), "Кому выдаётся", size=16, bold=True, color=GOLD, font="Georgia")
    add_bullets(s, Inches(7.4), Inches(2.9), Inches(5.0), Inches(3.3), [
        "Офицеры и прапорщики — ВКПО I категории",
        "Сержанты/солдаты по контракту — I, при отсутствии — II",
        "По призыву — в основном II категория",
        "На испытании контрактника — сезонные предметы II категории",
        "Крайний Север — допускается полная выдача по норме",
    ], size=13)
    footer(s, 9, total)

    # 10 Layers
    s = new_slide(prs)
    eyebrow(s, "Раздел 04")
    title(s, "Многослойный принцип ВКПО")
    if (ASSETS / "layers.png").exists():
        s.shapes.add_picture(str(ASSETS / "layers.png"), Inches(7.3), Inches(2.0), Inches(5.3), Inches(4.6))
    add_rect(s, Inches(0.7), Inches(2.0), Inches(6.3), Inches(4.6), fill=PANEL)
    add_text(s, Inches(0.95), Inches(2.2), Inches(5.8), Inches(0.4), "Типовые слои", size=16, bold=True, color=GOLD, font="Georgia")
    add_bullets(s, Inches(0.9), Inches(2.7), Inches(5.8), Inches(3.6), [
        "Базовый — футболка, трусы / бельё, флис, влагоотводящее",
        "Утепляющий — куртка и брюки утеплённые",
        "Основной полевой — костюм летний / демисезонный",
        "Защитный — костюм ветроводозащитный",
        "Головные уборы и обувь — фуражка, ушанка, балаклава, берцы",
        "Аксессуары — перчатки, рукавицы, шарф, баул",
        "Замена — по износу, не ранее срока, после сдачи",
    ], size=13)
    footer(s, 10, total)

    # 11 VKPO issuance
    s = new_slide(prs)
    eyebrow(s, "Раздел 04")
    title(s, "Особенности выдачи ВКПО")
    panel_card(s, Inches(0.7), Inches(2.1), Inches(6.0), Inches(4.5), "Призывники", [
        "Получают ВКПО преимущественно II категории",
        "При переводе сдают ВКПО, кроме отдельных предметов",
        "Исключения: фуражка летняя, футболка, трусы, костюм летний, ботинки для низких температур — с записью в аттестате",
        "Обязательны подгонка и клеймение",
    ])
    panel_card(s, Inches(6.95), Inches(2.1), Inches(5.7), Inches(4.5), "Контрактники", [
        "На Крайнем Севере — полная выдача по норме",
        "При переводе ВКПО/лётное/техническое/ССО — с записью в аттестате",
        "Проживающие вне казармы закрепляют инвентарь в регистрах учёта",
        "Отдельные предметы после срока могут перейти в собственность",
    ])
    footer(s, 11, total)

    # 12 Wear periods table
    s = new_slide(prs)
    eyebrow(s, "Раздел 05")
    title(s, "Сроки носки — типовые ориентиры", size=28)
    subtitle(s, "Точные сроки зависят от нормы, категории и климата. Сверяйте приказ № 500 и карточку довольствия.")
    rows = [
        ["Предмет", "Ориентир срока / нормы", "Примечание"],
        ["Футболки", "до 3 шт. на 1 год", "Личное / расходный характер"],
        ["Носки / портянки летние", "до 6 пар на 1 год", "Зимние — до 4 пар на 1 год"],
        ["Полуботинки", "1 пара на 2 года", "Зависит от нормы категории"],
        ["Полусапоги", "1 пара на 3 года", "Повседневная / специальная обувь"],
        ["Ремни поясной / брючный", "по 1 шт. на 5 лет", "Долгосрочные предметы"],
        ["Шапка зимняя (оф./прап.)", "1 шт. на 4 года", "п. 56.1 Порядка"],
        ["Балаклава", "1 шт. на 5 лет", "п. 56.1"],
        ["Футболка (оф./прап., норма № 1)", "2 шт. на 2 года", "п. 56.1"],
        ["Рукавицы утеплённые", "1 пара на 5 лет", "п. 56.1"],
        ["Сумка (баул)", "1 шт. на 10 лет", "п. 56.1"],
        ["Бельё влагоотводящее (контракт)", "2 комплекта на 1 год", "п. 56.2"],
    ]
    add_table(s, Inches(0.5), Inches(2.0), Inches(12.3), Inches(4.7), rows,
              col_widths=[Inches(4.2), Inches(4.0), Inches(4.1)])
    footer(s, 12, total)

    # 13 Climate adjustments
    s = new_slide(prs)
    eyebrow(s, "Раздел 05")
    title(s, "Климатические поправки к срокам")
    cards = [
        ("Умеренный климат", ["Базовый срок, указанный в норме снабжения", "Пример: костюм утеплённый — базовый срок"]),
        ("Холодный / особо холодный", ["Срок носки костюма утеплённого уменьшается на 1 год", "п. 62.1 Порядка — интенсивнее износ"]),
        ("Жаркий климат", ["Срок костюма утеплённого увеличивается на 1 год", "Для ряда предметов ВМФ — отдельные сроки (на 4 года)"]),
    ]
    for i, (h, lines) in enumerate(cards):
        panel_card(s, Inches(0.7 + i * 4.15), Inches(2.15), Inches(3.95), Inches(3.5), h, lines)
    add_rect(s, Inches(0.7), Inches(5.9), Inches(12), Inches(0.8), fill=PANEL2)
    add_text(s, Inches(0.95), Inches(6.05), Inches(11.5), Inches(0.5),
             "Срок нового предмета начинается со следующего дня после окончания срока ранее выданного одноимённого — «месяцы носки» не теряются.",
             size=12, color=SAND)
    footer(s, 13, total)

    # 14 Size system
    s = new_slide(prs)
    eyebrow(s, "Раздел 06")
    title(s, "Система размеров и ростовки")
    subtitle(s, "ГОСТ 23167-91 и ГОСТ 20881-91. Отпуск со складов — по заявленной ростовке.")
    if (ASSETS / "fitting.png").exists():
        s.shapes.add_picture(str(ASSETS / "fitting.png"), Inches(7.3), Inches(2.1), Inches(5.3), Inches(4.4))
    add_rect(s, Inches(0.7), Inches(2.1), Inches(6.3), Inches(4.4), fill=PANEL)
    add_text(s, Inches(0.95), Inches(2.3), Inches(5.8), Inches(0.4), "Определяющие признаки", size=16, bold=True, color=GOLD, font="Georgia")
    add_bullets(s, Inches(0.9), Inches(2.85), Inches(5.8), Inches(3.3), [
        "Мужчины: рост, обхват груди, обхват талии (полнота)",
        "Женщины: рост, обхват груди, обхват бёдер",
        "Обувь: длина стопы (метрическая, до 0,5 см)",
        "Головные уборы: обхват головы",
        "На сборных пунктах — запас ≈ 10% потребности призыва",
        "Данные → арматурная карточка → ростовка заявок",
    ], size=13)
    footer(s, 14, total)

    # 15 Growth / chest tables
    s = new_slide(prs)
    eyebrow(s, "Раздел 06")
    title(s, "Таблица ростов и размеров (мужчины)", size=28)
    rows1 = [
        ["Усл. рост", "Типовой, см", "Диапазон, см"],
        ["1", "158", "155 – 160,9"],
        ["2", "164", "161 – 166,9"],
        ["3", "170", "167 – 172,9"],
        ["4", "176", "173 – 178,9"],
        ["5", "182", "179 – 184,9"],
        ["6", "188", "185 – 191"],
    ]
    rows2 = [
        ["Размер", "Обхват груди", "Диапазон"],
        ["44", "88", "86–90"],
        ["46", "92", "90–94"],
        ["48", "96", "94–98"],
        ["50", "100", "98–102"],
        ["52", "104", "102–106"],
        ["54", "108", "106–110"],
        ["56", "112", "110–114"],
        ["58", "116", "114–118"],
        ["60", "120", "118–122"],
    ]
    add_text(s, Inches(0.7), Inches(1.95), Inches(5.5), Inches(0.35), "Росты", size=14, bold=True, color=GOLD)
    add_table(s, Inches(0.5), Inches(2.35), Inches(5.8), Inches(4.3), rows1,
              col_widths=[Inches(1.6), Inches(2.0), Inches(2.2)])
    add_text(s, Inches(6.9), Inches(1.95), Inches(5.5), Inches(0.35), "Размер по обхвату груди", size=14, bold=True, color=GOLD)
    add_table(s, Inches(6.7), Inches(2.35), Inches(6.0), Inches(4.3), rows2,
              col_widths=[Inches(1.6), Inches(2.1), Inches(2.3)])
    footer(s, 15, total)

    # 16 Fullness & measuring
    s = new_slide(prs)
    eyebrow(s, "Раздел 06")
    title(s, "Полнотные группы и снятие мерок")
    rows = [
        ["Обхват груди", "Обхват талии (гр. 1)"],
        ["88", "64 (62–66)"],
        ["96", "72 (70–74)"],
        ["104", "80 (78–82)"],
        ["112", "88 (86–90)"],
        ["120", "96 (94–98)"],
    ]
    add_table(s, Inches(0.5), Inches(2.1), Inches(5.5), Inches(3.8), rows,
              col_widths=[Inches(2.5), Inches(3.0)])
    add_rect(s, Inches(6.4), Inches(2.1), Inches(6.2), Inches(4.5), fill=PANEL)
    add_text(s, Inches(6.65), Inches(2.3), Inches(5.7), Inches(0.4), "Как снимают мерки", size=16, bold=True, color=GOLD, font="Georgia")
    add_bullets(s, Inches(6.6), Inches(2.85), Inches(5.7), Inches(3.4), [
        "Рост — от пола до верхушечной точки без обуви",
        "Обхват груди — горизонтально через сосковые точки",
        "Обхват талии / бёдер — по антропометрическим точкам",
        "Стопа — стопомером; размер в метрической системе",
        "Интервалы ГОСТ (оф./прап.): рост ±3, грудь ±2, талия ±3 см",
        "Инструменты: ростомер, сантиметр, стопомер",
    ], size=13)
    footer(s, 16, total)

    # 17 Climate zones
    s = new_slide(prs)
    eyebrow(s, "Раздел 07")
    title(s, "Климатические местности")
    panel_card(s, Inches(0.7), Inches(2.1), Inches(6.0), Inches(4.5), "Особо холодный — норма № 1", [
        "Части по перечню, утверждаемому НГШ ВС РФ",
        "Полевая форма — при наличии снежного покрова",
        "Дополнительное утепление",
        "Отдельные сроки: шапки, балаклавы, рукавицы",
        "Крайний Север — полная выдача ВКПО по норме",
    ])
    panel_card(s, Inches(6.95), Inches(2.1), Inches(5.7), Inches(4.5), "Жаркий климат и особые районы", [
        "Норма № 28 — загранформирования, ООН, жаркий климат",
        "ВМФ — облегчённые костюмы при плавании южнее 40° с. ш.",
        "Горная местность — отдельные правила специмущества",
        "Сроки утеплённых предметов корректируются",
    ])
    footer(s, 17, total)

    # 18 Special norms 1
    s = new_slide(prs)
    eyebrow(s, "Раздел 07")
    title(s, "Специальные нормы снабжения (часть 1)", size=28)
    rows = [
        ["Норма", "Категория обеспечиваемых"],
        ["№ 1", "Местности с особо холодным климатом"],
        ["№ 2", "Несение боевого дежурства"],
        ["№ 5", "Курсанты-лётчики (штурманы)"],
        ["№ 6", "РВСН и космические войска ВКС"],
        ["№ 8", "Совершающие прыжки с парашютом"],
        ["№ 9–11", "ВМФ, гидрография, аварийно-спасательные работы"],
        ["№ 12", "Постовая одежда"],
        ["№ 15–16", "Работы с РВ/ИИИ и спецбоеприпасами"],
        ["№ 17", "Штатные противопожарные подразделения"],
        ["№ 20", "Эксплуатация и обслуживание ВВТ"],
    ]
    add_table(s, Inches(0.7), Inches(2.0), Inches(12), Inches(4.7), rows,
              col_widths=[Inches(2.2), Inches(9.8)])
    footer(s, 18, total)

    # 19 Special norms 2
    s = new_slide(prs)
    eyebrow(s, "Раздел 07")
    title(s, "Специальные нормы снабжения (часть 2)", size=28)
    rows = [
        ["Норма", "Категория обеспечиваемых"],
        ["№ 21–25", "Медицина, санатории, роддома, переливание крови"],
        ["№ 27", "Творческий состав ансамблей и оркестров"],
        ["№ 28", "Загранформирования / жаркий климат"],
        ["№ 30", "Специальное и альпинистское имущество"],
        ["№ 32", "Спортивное имущество (изм. № 367/2024 — регби)"],
        ["№ 33", "Церемониальная форма почётного караула"],
        ["№ 34", "Силы специальных операций"],
        ["№ 36–37", "Палатки; брезенты и мягкие контейнеры"],
        ["№ 38–40", "Обозное и подковное имущество, мази"],
        ["№ 45", "Военный инновационный технополис «ЭРА»"],
    ]
    add_table(s, Inches(0.7), Inches(2.0), Inches(12), Inches(4.7), rows,
              col_widths=[Inches(2.2), Inches(9.8)])
    footer(s, 19, total)

    # 20 Issuance flow
    s = new_slide(prs)
    eyebrow(s, "Раздел 08")
    title(s, "Порядок получения и замены")
    steps = [
        ("1", "Прибытие", "Замеры, подгонка, первичный комплект, карточка"),
        ("2", "Носка", "Использование по назначению, закрепление инвентаря"),
        ("3", "Истечение срока", "Прибытие на склад / в ателье за заменой"),
        ("4", "Сдача", "Изношенное сдаётся — затем выдаётся новое"),
        ("5", "Учёт", "Карточка, аттестат, акты преждевременного износа"),
    ]
    for i, (n, h, d) in enumerate(steps):
        l = Inches(0.45 + i * 2.55)
        add_rect(s, l, Inches(2.4), Inches(2.4), Inches(3.6), fill=PANEL)
        add_text(s, l + Inches(0.15), Inches(2.6), Inches(2.1), Inches(0.5), n, size=24, bold=True, color=OLIVE_B, font="Georgia")
        add_text(s, l + Inches(0.15), Inches(3.25), Inches(2.1), Inches(0.6), h, size=14, bold=True, color=SAND)
        add_text(s, l + Inches(0.15), Inches(4.0), Inches(2.1), Inches(1.6), d, size=12, color=MUTED)
    footer(s, 20, total)

    # 21 Warehouse
    s = new_slide(prs)
    eyebrow(s, "Раздел 08")
    title(s, "Складская логистика")
    subtitle(s, "Отпуск с ЦМТО — по нарядам вещевой службы округа по ростовке.")
    if (ASSETS / "warehouse.png").exists():
        s.shapes.add_picture(str(ASSETS / "warehouse.png"), Inches(0.7), Inches(2.15), Inches(6.2), Inches(4.4))
    add_rect(s, Inches(7.2), Inches(2.15), Inches(5.4), Inches(4.4), fill=PANEL)
    add_text(s, Inches(7.45), Inches(2.35), Inches(5.0), Inches(0.4), "В обеспеченности учитывают", size=15, bold=True, color=GOLD, font="Georgia")
    add_bullets(s, Inches(7.4), Inches(2.9), Inches(5.0), Inches(3.3), [
        "Всё новое имущество",
        "Б/у, срок которого не истекает в периоде",
        "Инвентарное, выслужившее срок, но годное",
        "Ростовка = % предметов разных размеров",
        "Подменный фонд — из сданного годного имущества",
        "Рабочая одежда призывникам — до 100% численности",
    ], size=13)
    footer(s, 21, total)

    # 22 Documents
    s = new_slide(prs)
    eyebrow(s, "Раздел 08")
    title(s, "Документы и учёт")
    cards = [
        ("Карточка вещевого довольствия", ["Основной личный учёт выдач и сроков", "Первая точка сверки «что положено»"]),
        ("Аттестат военнослужащего", ["При переводе фиксирует следующее имущество", "ВКПО, лётное, техническое, ССО, исключения"]),
        ("Наряды / акты", ["Складской отпуск по ростовке", "Акты при преждевременном износе и списании"]),
    ]
    for i, (h, lines) in enumerate(cards):
        panel_card(s, Inches(0.7 + i * 4.15), Inches(2.15), Inches(3.95), Inches(3.3), h, lines)
    add_rect(s, Inches(0.7), Inches(5.7), Inches(12), Inches(0.95), fill=PANEL2)
    add_text(s, Inches(0.95), Inches(5.9), Inches(11.5), Inches(0.6),
             "При увольнении к дню исключения из списков инвентарь сдаётся; военнослужащий обеспечивается положенным имуществом личного пользования.",
             size=12, color=SAND)
    footer(s, 22, total)

    # 23 Compensation
    s = new_slide(prs)
    eyebrow(s, "Раздел 08")
    title(s, "Денежная компенсация вместо вещевого")
    subtitle(s, "Основание — ст. 14 Закона № 76-ФЗ и Правила к ПП № 390.")
    panel_card(s, Inches(0.7), Inches(2.15), Inches(6.0), Inches(4.4), "Кому доступно", [
        "Отдельным категориям контрактников",
        "Вместо предметов личного пользования (не любого инвентаря)",
        "Обычно после истечения срока носки",
        "При увольнении — за неполученное (по правилам)",
        "Размеры — по решениям Правительства",
    ])
    panel_card(s, Inches(6.95), Inches(2.15), Inches(5.7), Inches(4.4), "Алгоритм", [
        "Проверить право категории по Правилам к ПП № 390",
        "Подать рапорт на имя командира",
        "Справка вещевой службы о неполучении / сроках",
        "Финансовая служба производит выплату",
        "Стоимость — в т. ч. по Распоряжению № 1014-р",
    ])
    footer(s, 23, total)

    # 24 Bath laundry
    s = new_slide(prs)
    eyebrow(s, "Раздел 08")
    title(s, "Банно-прачечное обслуживание")
    cards = [
        ("Что входит", ["Помывка личного состава", "Стирка белья и обмундирования", "Банный инвентарь и моющие", "Ремонт имущества вещслужбы"]),
        ("Нормы расхода", ["В приложениях к приказу № 500", "Банный инвентарь, моющие, ремматериалы", "Средства ухода за обувью", "Оборудование прачечных"]),
        ("Постельные", ["Инвентарь подразделения / коечного фонда", "Полотенца для ног — из выслуживших срок, но годных", "Учёт и нормы содержания на койку"]),
    ]
    for i, (h, lines) in enumerate(cards):
        panel_card(s, Inches(0.7 + i * 4.15), Inches(2.15), Inches(3.95), Inches(4.4), h, lines)
    footer(s, 24, total)

    # 25 Updates 2024-2025
    s = new_slide(prs)
    eyebrow(s, "Актуально")
    title(s, "Изменения 2024–2025")
    panel_card(s, Inches(0.7), Inches(2.15), Inches(6.0), Inches(4.4), "Приказ МО № 367 от 27.06.2024", [
        "Изменения норм снабжения к приказу № 500",
        "Специальное, санитарно-хозяйственное, спортивное имущество",
        "Зарегистрирован в Минюсте 07.08.2024 № 79039",
        "Пример: дополнение нормы № 32 — экипировка для регби-7",
    ])
    panel_card(s, Inches(6.95), Inches(2.15), Inches(5.7), Inches(4.4), "ПП РФ № 1116 от 28.07.2025", [
        "Изменения в постановление № 390",
        "Приостановка до 31.12.2030 выдачи отдельных предметов",
        "Перечень — в постановлении (в т. ч. отдельные предметы авиации/ПВО)",
        "Перед планированием сверяйте актуальную редакцию ПП № 390",
    ])
    footer(s, 25, total)

    # 26 Checklist
    s = new_slide(prs)
    eyebrow(s, "Практика")
    title(s, "Чек-лист военнослужащего")
    panel_card(s, Inches(0.7), Inches(2.15), Inches(6.0), Inches(4.4), "При прибытии / контракте", [
        "Пройти обмер и примерку",
        "Получить комплект по норме",
        "Проверить записи в карточке",
        "Уточнить климатическую норму части",
        "Сохранить сведения о сроках носки",
    ])
    panel_card(s, Inches(6.95), Inches(2.15), Inches(5.7), Inches(4.4), "В ходе службы", [
        "Менять имущество по срокам, не выбрасывая инвентарь",
        "При преждевременном износе по службе — акт",
        "При переводе — корректный аттестат",
        "При праве на компенсацию — рапорт + справка",
        "Перед увольнением — сдать инвентарь",
    ])
    footer(s, 26, total)

    # 27 Sources
    s = new_slide(prs)
    eyebrow(s, "Справки")
    title(s, "Основные источники")
    subtitle(s, "Презентация обзорная; юридически значимы полные тексты в актуальной редакции.")
    sources = [
        ("76-ФЗ", "Федеральный закон «О статусе военнослужащих», ст. 14"),
        ("ПП № 390", "Постановление Правительства РФ от 22.06.2006 № 390 (изм. ПП № 1116/2025)"),
        ("МО № 500", "Приказ Министра обороны РФ от 14.08.2017 № 500 (ред. от 27.06.2024)"),
        ("МО № 367", "Приказ Министра обороны РФ от 27.06.2024 № 367"),
        ("ГОСТ", "ГОСТ 23167-91, ГОСТ 20881-91 — типовые фигуры и шкалы размеров"),
        ("1014-р", "Распоряжение Правительства РФ от 25.05.2016 № 1014-р"),
    ]
    for i, (code, desc) in enumerate(sources):
        t = Inches(2.05 + i * 0.75)
        add_rect(s, Inches(0.7), t, Inches(12), Inches(0.68), fill=PANEL)
        add_text(s, Inches(0.95), t + Inches(0.15), Inches(2.2), Inches(0.4), code, size=13, bold=True, color=GOLD, font="Georgia")
        add_text(s, Inches(3.3), t + Inches(0.15), Inches(9.1), Inches(0.4), desc, size=13, color=MUTED)
    footer(s, 27, total)

    # 28 Closing
    s = new_slide(prs)
    add_text(s, Inches(0.8), Inches(2.0), Inches(11.5), Inches(0.4),
             "ВЕЩЕВАЯ СЛУЖБА ВС РФ", size=14, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
    add_text(s, Inches(0.8), Inches(2.6), Inches(11.5), Inches(1.4),
             "Готовность начинается\nс обеспечения", size=40, bold=True, color=SAND, font="Georgia", align=PP_ALIGN.CENTER)
    add_text(s, Inches(1.5), Inches(4.3), Inches(10.3), Inches(1.2),
             "Норма → ростовка → выдача → носка → замена.\n"
             "Полный цикл держится на приказе № 500, постановлении № 390 и дисциплине учёта в воинской части.",
             size=15, color=MUTED, align=PP_ALIGN.CENTER)
    footer(s, 28, total)

    prs.save(str(OUT))
    print(f"Saved: {OUT}")
    print(f"Size: {OUT.stat().st_size / 1024 / 1024:.1f} MB")


if __name__ == "__main__":
    build()
