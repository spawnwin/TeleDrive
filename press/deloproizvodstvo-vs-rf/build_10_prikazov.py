#!/usr/bin/env python3
"""10 учебных образцов приказов по ИД-2017 (приказ МО РФ № 170)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
DOCX_DIR = ROOT / "docx" / "obrazcy_10"
IMG_DIR = ROOT / "images" / "obrazcy_10"
DOCX_DIR.mkdir(parents=True, exist_ok=True)
IMG_DIR.mkdir(parents=True, exist_ok=True)

FONT_REG = "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf"
FONT_SANS = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
FONT_SANS_BOLD = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"

# 10 типовых приказов по вопросам из п. 56 ИД-2017 / повседневной деятельности
ORDERS = [
    {
        "file": "01_sutochnyy_naryad",
        "number": "121/НР",
        "date": "«05» августа 2026 г.",
        "title": "О составе суточного наряда\nвойсковой части 00000\nна 6 августа 2026 г.",
        "preamble": (
            "В соответствии с Уставом внутренней службы Вооруженных Сил Российской Федерации "
            "и в целях поддержания внутреннего порядка, организации охраны и обороны объектов "
            "войсковой части 00000,"
        ),
        "points": [
            "Назначить суточный наряд войсковой части 00000 на 6 августа 2026 г. в составе согласно расчету (приложение к настоящему приказу).",
            "Дежурному по воинской части капитану Смирнову А.В. принять дела и должность от предыдущего дежурного к 17.00 5 августа 2026 г.",
            "Командирам подразделений к 16.00 5 августа 2026 г. представить личный состав суточного наряда на строевой плац.",
            "Начальнику штаба организовать инструктаж суточного наряда.",
            "Контроль за выполнением настоящего приказа возложить на начальника штаба войсковой части 00000.",
        ],
        "appendix": True,
        "appendix_title": "РАСЧЕТ\nсуточного наряда войсковой части 00000\nна 6 августа 2026 г.",
        "appendix_points": [
            "Дежурный по воинской части — капитан Смирнов А.В.",
            "Помощник дежурного по воинской части — старший лейтенант Козлов Д.С.",
            "Дежурный по парку — старший прапорщик Волков Н.И.",
            "Дневальные по ротам — по отдельному списку командиров подразделений.",
        ],
    },
    {
        "file": "02_zachislenie_v_spiski",
        "number": "122",
        "date": "«05» августа 2026 г.",
        "title": "О зачислении в списки\nличного состава войсковой части 00000",
        "preamble": (
            "На основании предписания командира войсковой части 11111 от 1 августа 2026 г. № 45 "
            "и в соответствии с порядком ведения учета личного состава,"
        ),
        "points": [
            "Зачислить в списки личного состава войсковой части 00000 и поставить на все виды довольствия с 5 августа 2026 г. рядового Кузнецова Ивана Петровича, 2004 года рождения.",
            "Назначить рядового Кузнецова И.П. на воинскую должность стрелка 1 мотострелковой роты.",
            "Начальнику строевого отделения внести изменения в учетные документы в срок до 7 августа 2026 г.",
            "Контроль за выполнением настоящего приказа возложить на начальника штаба войсковой части 00000.",
        ],
        "appendix": False,
    },
    {
        "file": "03_komandirovka",
        "number": "123",
        "date": "«06» августа 2026 г.",
        "title": "О направлении в служебную командировку",
        "preamble": (
            "В целях обеспечения участия в учебно-методическом сборе по вопросам организации "
            "делопроизводства и в соответствии с планом мероприятий на август 2026 г.,"
        ),
        "points": [
            "Направить в служебную командировку в г. Москву с 10 по 14 августа 2026 г. начальника строевого отделения капитана Иванова И.И. для участия в учебно-методическом сборе.",
            "Цель командировки — изучение порядка применения Инструкции по делопроизводству в Вооруженных Силах Российской Федерации (ИД-2017).",
            "Начальнику финансовой службы обеспечить выплату командировочных расходов в установленном порядке.",
            "Капитану Иванову И.И. по возвращении из командировки в трехдневный срок представить письменный отчет.",
            "Контроль за выполнением настоящего приказа оставляю за собой.",
        ],
        "appendix": False,
    },
    {
        "file": "04_otpusk",
        "number": "124",
        "date": "«06» августа 2026 г.",
        "title": "О предоставлении отпуска",
        "preamble": (
            "В соответствии с Федеральным законом от 27 мая 1998 г. № 76-ФЗ «О статусе военнослужащих» "
            "и на основании рапорта майора Орлова С.А.,"
        ),
        "points": [
            "Предоставить майору Орлову Сергею Александровичу основной отпуск за 2026 год продолжительностью 30 суток с 15 августа по 13 сентября 2026 г. включительно, не считая времени на путь следования.",
            "На время отпуска исполнение обязанностей по должности возложить на капитана Белова Р.К.",
            "Начальнику строевого отделения оформить отпускной билет и внести сведения в учетные документы.",
            "Контроль за выполнением настоящего приказа возложить на начальника штаба войсковой части 00000.",
        ],
        "appendix": False,
    },
    {
        "file": "05_pooshchrenie",
        "number": "125",
        "date": "«07» августа 2026 г.",
        "title": "О поощрении личного состава",
        "preamble": (
            "За добросовестное исполнение должностных обязанностей, высокие результаты "
            "в боевой подготовке и в связи с успешным проведением контрольной проверки,"
        ),
        "points": [
            "Объявить благодарность старшему лейтенанту Морозову П.А.",
            "Наградить грамотой командира войсковой части сержанта Лебедева К.В.",
            "Выдать ценный подарок прапорщику Соколову М.Н.",
            "Настоящий приказ довести до всего личного состава войсковой части в части, касающейся.",
            "Контроль за выполнением настоящего приказа возложить на заместителя командира войсковой части 00000 по военно-политической работе.",
        ],
        "appendix": False,
    },
    {
        "file": "06_naznachenie_na_dolzhnost",
        "number": "126",
        "date": "«07» августа 2026 г.",
        "title": "О назначении на воинскую должность",
        "preamble": (
            "В соответствии с полномочиями командира воинской части по назначению военнослужащих, "
            "проходящих военную службу по призыву, на воинские должности и на основании представления "
            "командира 2 мотострелковой роты,"
        ),
        "points": [
            "Освободить ефрейтора Николаева А.С. от воинской должности стрелка 2 мотострелковой роты.",
            "Назначить ефрейтора Николаева А.С. на воинскую должность командира отделения 2 мотострелковой роты с 8 августа 2026 г.",
            "Начальнику строевого отделения внести изменения в штатно-должностную книгу и учетные документы в срок до 9 августа 2026 г.",
            "Контроль за выполнением настоящего приказа возложить на начальника штаба войсковой части 00000.",
        ],
        "appendix": False,
    },
    {
        "file": "07_boevaya_podgotovka",
        "number": "127",
        "date": "«08» августа 2026 г.",
        "title": "Об организации боевой подготовки\nв августе 2026 г.",
        "preamble": (
            "В целях качественной организации боевой подготовки личного состава "
            "и выполнения плана боевой подготовки на 2026 учебный год,"
        ),
        "points": [
            "Утвердить План боевой подготовки войсковой части 00000 на август 2026 г. (приложение к настоящему приказу).",
            "Командирам подразделений организовать проведение занятий в соответствии с утвержденным планом и обеспечить явку личного состава.",
            "Заместителю командира по военно-политической работе обеспечить военно-политическое сопровождение мероприятий боевой подготовки.",
            "Начальнику штаба еженедельно докладывать о ходе выполнения плана боевой подготовки.",
            "Контроль за выполнением настоящего приказа оставляю за собой.",
        ],
        "appendix": True,
        "appendix_title": "ПЛАН\nбоевой подготовки войсковой части 00000\nна август 2026 г.",
        "appendix_points": [
            "Огневая подготовка — по средам, строевой плац / тир.",
            "Тактическая подготовка — по вторникам и четвергам.",
            "Физическая подготовка — ежедневно, утренние часы.",
            "Техническая подготовка — по пятницам, парк боевых машин.",
            "Контрольные занятия — 28–29 августа 2026 г.",
        ],
    },
    {
        "file": "08_inventarizaciya",
        "number": "128",
        "date": "«08» августа 2026 г.",
        "title": "О проведении инвентаризации\nматериальных средств",
        "preamble": (
            "В целях обеспечения сохранности материальных средств, проверки наличия и состояния "
            "имущества и в соответствии с порядком проведения инвентаризации в Вооруженных Силах "
            "Российской Федерации,"
        ),
        "points": [
            "Провести инвентаризацию материальных средств складов вещевой и продовольственной служб в период с 15 по 25 августа 2026 г.",
            "Назначить инвентаризационную комиссию в составе: председатель — майор Федоров В.Л.; члены комиссии — капитан Егоров И.Н., старший лейтенант Павлов С.К.",
            "Комиссии представить акт инвентаризации на утверждение к 28 августа 2026 г.",
            "Материально ответственным лицам к началу инвентаризации завершить оформление приходно-расходных документов.",
            "Контроль за выполнением настоящего приказа возложить на заместителя командира войсковой части 00000 по вооружению.",
        ],
        "appendix": False,
    },
    {
        "file": "09_sohrannost_oruzhiya",
        "number": "129",
        "date": "«09» августа 2026 г.",
        "title": "Об обеспечении сохранности оружия,\nбоеприпасов и взрывчатых веществ",
        "preamble": (
            "В целях усиления контроля за сохранностью оружия, боеприпасов и взрывчатых веществ "
            "и исключения предпосылок к их утрате,"
        ),
        "points": [
            "Командирам подразделений до 12 августа 2026 г. провести проверку наличия, состояния учета и условий хранения оружия и боеприпасов.",
            "Начальнику службы ракетно-артиллерийского вооружения организовать внезапную проверку комнат для хранения оружия не реже двух раз в неделю.",
            "Дежурному по воинской части при приеме (сдаче) дежурства проверять целостность печатей и наличие ключей от комнат хранения оружия.",
            "Обо всех нарушениях немедленно докладывать командиру войсковой части.",
            "Контроль за выполнением настоящего приказа возложить на начальника штаба войсковой части 00000.",
        ],
        "appendix": False,
    },
    {
        "file": "10_pozharnaya_bezopasnost",
        "number": "130",
        "date": "«09» августа 2026 г.",
        "title": "Об усилении мер пожарной безопасности",
        "preamble": (
            "В целях предупреждения пожаров, обеспечения пожарной безопасности объектов "
            "войсковой части 00000 и во исполнение требований нормативных правовых актов "
            "по пожарной безопасности,"
        ),
        "points": [
            "Утвердить Инструкцию о мерах пожарной безопасности в войсковой части 00000 (приложение к настоящему приказу).",
            "Назначить ответственным за пожарную безопасность в войсковой части 00000 майора Григорьева А.Н.",
            "Командирам подразделений до 15 августа 2026 г. провести занятия с личным составом по мерам пожарной безопасности и оформить листы ознакомления.",
            "Начальнику инженерной службы проверить наличие и исправность средств пожаротушения к 14 августа 2026 г.",
            "Контроль за выполнением настоящего приказа возложить на начальника штаба войсковой части 00000.",
        ],
        "appendix": True,
        "appendix_title": "ИНСТРУКЦИЯ\nо мерах пожарной безопасности\nв войсковой части 00000",
        "appendix_points": [
            "На территории воинской части запрещается разведение костров вне установленных мест.",
            "Курение допускается только в специально отведенных местах.",
            "Подъезды к источникам воды и средствам пожаротушения должны быть свободны.",
            "О каждом случае задымления или возгорания немедленно сообщать дежурному по воинской части.",
            "Ответственные лица ежедневно проверяют исправность пожарного инвентаря.",
        ],
    },
]


def set_run_font(run, size=14, bold=False, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_paragraph_format(
    p,
    *,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    space_before=0,
    space_after=0,
    line_spacing=1.15,
    first_line=None,
):
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if first_line is None:
        pf.first_line_indent = Cm(0)
    else:
        pf.first_line_indent = Cm(first_line)


def add_centered(doc, text, *, size=14, bold=False, space_before=0, space_after=0, line_spacing=1.0):
    p = doc.add_paragraph()
    set_paragraph_format(
        p,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=space_before,
        space_after=space_after,
        line_spacing=line_spacing,
    )
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_body(doc, text, *, first_line=1.25, space_after=4):
    p = doc.add_paragraph()
    set_paragraph_format(
        p,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        space_after=space_after,
        line_spacing=1.15,
        first_line=first_line,
    )
    run = p.add_run(text)
    set_run_font(run, size=14)
    return p


def add_empty(doc, count=1):
    for _ in range(count):
        p = doc.add_paragraph()
        set_paragraph_format(p, space_before=0, space_after=0, line_spacing=1.0)
        set_run_font(p.add_run(""), size=14)


def set_section_a4_margins(section):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(30)
    section.right_margin = Mm(15)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)


def add_signatures(doc):
    add_empty(doc, 2)
    table = doc.add_table(rows=2, cols=2)
    cells = [
        ("Командир войсковой части 00000\nполковник", "И. Сидоров"),
        ("Начальник штаба\nвойсковой части 00000\nподполковник", "П. Петров"),
    ]
    for row_idx, (left, right) in enumerate(cells):
        cell_l = table.rows[row_idx].cells[0]
        cell_r = table.rows[row_idx].cells[1]
        cell_l.text = ""
        cell_r.text = ""
        p_l = cell_l.paragraphs[0]
        set_paragraph_format(p_l, align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.0)
        for j, line in enumerate(left.split("\n")):
            if j:
                p_l.add_run("\n")
            set_run_font(p_l.add_run(line), size=14)
        p_r = cell_r.paragraphs[0]
        set_paragraph_format(p_r, align=WD_ALIGN_PARAGRAPH.RIGHT, line_spacing=1.0)
        set_run_font(p_r.add_run("____________  " + right), size=14)
        if row_idx == 0:
            # spacer row effect via empty paragraph after table is enough
            pass
    add_empty(doc, 1)


def build_order_docx(order: dict) -> Path:
    doc = Document()
    set_section_a4_margins(doc.sections[0])
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    add_centered(doc, "МИНИСТЕРСТВО ОБОРОНЫ РОССИЙСКОЙ ФЕДЕРАЦИИ", size=12, bold=True)
    add_centered(doc, "(Минобороны России)", size=12, space_after=2)
    add_centered(doc, "ВОЙСКОВАЯ ЧАСТЬ 00000", size=14, bold=True, space_after=6)
    add_centered(doc, "ПРИКАЗ", size=16, bold=True, space_before=6, space_after=8)

    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.0)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT)
    set_run_font(p.add_run(order["date"]), size=14)
    set_run_font(p.add_run(f"\t№ {order['number']}"), size=14)

    add_centered(doc, "г. Москва", size=14, space_before=2, space_after=10)
    add_centered(doc, order["title"], size=14, bold=True, space_before=6, space_after=10, line_spacing=1.0)
    add_body(doc, order["preamble"], space_after=6)
    add_centered(doc, "ПРИКАЗЫВАЮ:", size=14, bold=True, space_before=4, space_after=8)

    for i, text in enumerate(order["points"], 1):
        add_body(doc, f"{i}. {text}")

    add_signatures(doc)

    if order.get("appendix"):
        doc.add_page_break()
        p = doc.add_paragraph()
        set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.RIGHT, line_spacing=1.0)
        p.paragraph_format.left_indent = Cm(7)
        mark = (
            f"Приложение\nк приказу командира\nвойсковой части 00000\n"
            f"от {order['date'].replace('«', '').replace('»', '')} № {order['number']}"
        )
        for j, line in enumerate(mark.split("\n")):
            if j:
                p.add_run("\n")
            set_run_font(p.add_run(line), size=12)

        add_centered(
            doc,
            order["appendix_title"],
            size=14,
            bold=True,
            space_before=18,
            space_after=12,
            line_spacing=1.0,
        )
        for i, text in enumerate(order.get("appendix_points", []), 1):
            add_body(doc, f"{i}. {text}")
        add_empty(doc, 1)
        p = doc.add_paragraph()
        set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.0)
        set_run_font(p.add_run("―" * 12), size=14)

    out = DOCX_DIR / f"{order['file']}.docx"
    doc.save(out)
    return out


def font(path, size):
    return ImageFont.truetype(path, size)


def wrap_lines(draw, text, fnt, max_width):
    words = text.split()
    lines, cur = [], ""
    for w in words:
        test = (cur + " " + w).strip()
        if draw.textbbox((0, 0), test, font=fnt)[2] <= max_width:
            cur = test
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def build_order_preview(order: dict, index: int) -> Path:
    W, H = 900, 1270
    img = Image.new("RGB", (W, H), (255, 255, 255))
    draw = ImageDraw.Draw(img)
    mm = 150 / 25.4
    L, R, T = int(28 * mm), int(14 * mm), int(18 * mm)

    f12b = font(FONT_BOLD, 18)
    f14 = font(FONT_REG, 20)
    f14b = font(FONT_BOLD, 20)
    f16b = font(FONT_BOLD, 26)
    f_small = font(FONT_SANS, 14)

    cx = (L + W - R) // 2
    y = T

    def center(text, yy, fnt, gap=3):
        bb = draw.textbbox((0, 0), text, font=fnt)
        draw.text((cx - (bb[2] - bb[0]) // 2, yy), text, fill=(0, 0, 0), font=fnt)
        return yy + (bb[3] - bb[1]) + gap

    # badge
    draw.rounded_rectangle((W - 160, 16, W - 20, 52), radius=8, fill=(180, 50, 50))
    draw.text((W - 148, 24), f"Образец {index:02d}", fill=(255, 255, 255), font=f_small)

    y = center("МИНИСТЕРСТВО ОБОРОНЫ", y, f12b)
    y = center("РОССИЙСКОЙ ФЕДЕРАЦИИ", y, f12b)
    y = center("ВОЙСКОВАЯ ЧАСТЬ 00000", y + 4, f14b)
    y = center("ПРИКАЗ", y + 10, f16b)
    y += 12
    draw.text((L, y), order["date"], fill=(0, 0, 0), font=f14)
    num = f"№ {order['number']}"
    nb = draw.textbbox((0, 0), num, font=f14)
    draw.text((W - R - (nb[2] - nb[0]), y), num, fill=(0, 0, 0), font=f14)
    y += 28
    y = center("г. Москва", y, f14)
    y += 12
    for line in order["title"].split("\n"):
        y = center(line, y, f14b, gap=2)
    y += 12

    indent = int(10 * mm)
    for i, line in enumerate(wrap_lines(draw, order["preamble"], f14, W - L - R - 8)):
        x = L + (indent if i == 0 else 0)
        draw.text((x, y), line, fill=(0, 0, 0), font=f14)
        y += 26
        if y > H - 220:
            break
    y += 8
    y = center("ПРИКАЗЫВАЮ:", y, f14b)
    y += 8
    for n, pt in enumerate(order["points"], 1):
        lines = wrap_lines(draw, f"{n}. {pt}", f14, W - L - R - indent - 8)
        for i, line in enumerate(lines):
            draw.text((L + (indent if i == 0 else 0), y), line, fill=(0, 0, 0), font=f14)
            y += 24
            if y > H - 160:
                break
        y += 4
        if y > H - 160:
            break

    y = max(y + 20, H - 140)
    draw.multiline_text((L, y), "Командир в/ч 00000\nполковник", fill=(0, 0, 0), font=f14, spacing=2)
    sig = "______ И. Сидоров"
    sb = draw.textbbox((0, 0), sig, font=f14)
    draw.text((W - R - (sb[2] - sb[0]), y + 14), sig, fill=(0, 0, 0), font=f14)

    draw.rectangle((0, H - 32, W, H), fill=(245, 245, 245))
    draw.text((12, H - 24), "Учебный образец по ИД-2017 (приказ МО РФ № 170)", fill=(90, 90, 90), font=f_small)

    out = IMG_DIR / f"{order['file']}.png"
    img.save(out, "PNG")
    return out


def build_index_docx() -> Path:
    doc = Document()
    set_section_a4_margins(doc.sections[0])
    add_centered(doc, "ПЕРЕЧЕНЬ\nучебных образцов приказов\nпо ИД-2017", size=14, bold=True, space_after=12, line_spacing=1.0)
    add_body(
        doc,
        "Ниже приведен комплект из 10 учебных образцов приказов командира войсковой части, "
        "оформленных по Инструкции по делопроизводству в Вооруженных Силах Российской Федерации "
        "(ИД-2017), утвержденной приказом Министра обороны Российской Федерации от 4 апреля 2017 г. № 170.",
        space_after=8,
    )
    for i, order in enumerate(ORDERS, 1):
        title = order["title"].replace("\n", " ")
        add_body(
            doc,
            f"{i}. № {order['number']} от {order['date']} — {title}. Файл: {order['file']}.docx",
            space_after=3,
        )
    add_empty(doc, 1)
    add_body(
        doc,
        "Все образцы имеют единые поля (Л 30 / П 15 / В 20 / Н 20 мм), шрифт Times New Roman 14, "
        "абзацный отступ 1,25 см, структуру «констатация → ПРИКАЗЫВАЮ → пункты → подписи».",
        space_after=0,
    )
    out = DOCX_DIR / "00_perechen_obrazcov.docx"
    doc.save(out)
    return out


def build_montage() -> Path:
    # 2x5 grid of previews
    thumbs = []
    for i, order in enumerate(ORDERS, 1):
        path = IMG_DIR / f"{order['file']}.png"
        im = Image.open(path).convert("RGB")
        im.thumbnail((420, 594))
        thumbs.append(im)

    cols, rows = 5, 2
    pad = 16
    tw, th = thumbs[0].size
    W = cols * tw + (cols + 1) * pad
    H = rows * th + (rows + 1) * pad + 60
    canvas = Image.new("RGB", (W, H), (240, 240, 238))
    draw = ImageDraw.Draw(canvas)
    draw.text((pad, 18), "10 учебных образцов приказов по ИД-2017", fill=(30, 30, 30), font=font(FONT_SANS_BOLD, 28))
    for idx, im in enumerate(thumbs):
        r, c = divmod(idx, cols)
        # wait - 5 cols, 2 rows: idx 0..4 row0, 5..9 row1
        r = 0 if idx < 5 else 1
        c = idx if idx < 5 else idx - 5
        x = pad + c * (tw + pad)
        y = 60 + pad + r * (th + pad)
        canvas.paste(im, (x, y))
        draw.rectangle((x, y, x + im.width - 1, y + im.height - 1), outline=(180, 50, 50), width=2)
    out = ROOT / "images" / "06_desyat_obrazcov_montage.png"
    canvas.save(out, "PNG")
    return out


def main():
    files = [build_index_docx()]
    for i, order in enumerate(ORDERS, 1):
        files.append(build_order_docx(order))
        files.append(build_order_preview(order, i))
    files.append(build_montage())
    for f in files:
        print("OK:", f)


if __name__ == "__main__":
    main()
